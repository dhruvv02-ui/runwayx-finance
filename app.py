import streamlit as st
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import io
import os
import time
import base64
from reconciliation_agent import AIFinanceController
from generate_data import generate_synthetic_data

_favicon_svg = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 110 100">
    <polygon points="58,6 74,6 68,16 52,16" fill="#00E5FF" />
    <polygon points="35,16 55,16 35,50 44,50 41,58 30,58 12,90 2,90 22,58 17,58 20,50 26,50" fill="#0E1E38" />
    <polygon points="63,16 79,16 58,50 67,50 64,58 53,58 35,90 25,90 45,58 40,58 43,50 49,50" fill="#2563EB" />
</svg>"""

runwayx_favicon = f"data:image/svg+xml;base64,{base64.b64encode(_favicon_svg.encode()).decode()}"

st.set_page_config(
    page_title="RUNWAYX | Autonomous Reconciliation", 
    page_icon=runwayx_favicon, 
    layout="wide", 
    initial_sidebar_state="expanded"
)

st.markdown("""
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&family=Space+Grotesk:wght@500;600;700;800&display=swap" rel="stylesheet">
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Material+Symbols+Rounded:opsz,wght,FILL,GRAD@24,400,0,0" />
""", unsafe_allow_html=True)

if "has_loaded" not in st.session_state:
    st.session_state["has_loaded"] = True
    show_splash = True
else:
    show_splash = False

if show_splash:
    st.markdown("""
    <style>
        @keyframes pulseLogoGlow {
            0%, 100% { transform: scale(1); filter: drop-shadow(0 0 16px rgba(37, 99, 235, 0.45)); }
            50% { transform: scale(1.04); filter: drop-shadow(0 0 32px rgba(0, 229, 255, 0.7)); }
        }
        @keyframes progressSweep {
            0% { width: 0%; }
            50% { width: 65%; }
            100% { width: 100%; }
        }
        @keyframes cuteBounceDot {
            0%, 100% { transform: translateY(0px) scale(1); box-shadow: 0 0 8px rgba(0, 229, 255, 0.8); }
            50% { transform: translateY(-5px) scale(1.3); box-shadow: 0 0 14px rgba(0, 229, 255, 1); }
        }
        @keyframes fadeOutLoader {
            0% { opacity: 1; visibility: visible; }
            82% { opacity: 1; visibility: visible; }
            100% { opacity: 0; visibility: hidden; pointer-events: none; }
        }

        #splash-screen {
            position: fixed;
            top: 0;
            left: 0;
            width: 100vw;
            height: 100vh;
            background-color: #020611;
            background-image: 
                radial-gradient(circle at 50% 45%, rgba(37, 99, 235, 0.18) 0%, transparent 55%),
                radial-gradient(circle at 50% 50%, rgba(2, 132, 199, 0.15) 0%, transparent 65%),
                linear-gradient(180deg, #020611 0%, #030816 100%);
            z-index: 9999999;
            display: flex;
            flex-direction: column;
            justify-content: center;
            align-items: center;
            animation: fadeOutLoader 2.2s cubic-bezier(0.77, 0, 0.175, 1) forwards;
            pointer-events: none;
        }

        .splash-logo-container {
            position: relative;
            display: flex;
            justify-content: center;
            align-items: center;
            margin-bottom: 20px;
            animation: pulseLogoGlow 2s ease-in-out infinite;
        }

        .splash-brand {
            font-family: 'Space Grotesk', sans-serif;
            font-weight: 800;
            font-size: 36px;
            font-style: italic;
            letter-spacing: -0.5px;
            color: #FFFFFF;
            margin-bottom: 4px;
        }

        .splash-subtext {
            font-family: 'Plus Jakarta Sans', sans-serif;
            font-size: 11px;
            font-weight: 700;
            letter-spacing: 3.5px;
            text-transform: uppercase;
            color: #94A3B8;
            margin-bottom: 22px;
        }

        .splash-progress-wrapper {
            display: flex;
            align-items: center;
            gap: 12px;
        }

        .splash-progress-track {
            width: 210px;
            height: 4px;
            background: rgba(255, 255, 255, 0.08);
            border-radius: 9999px;
            overflow: hidden;
            position: relative;
        }

        .splash-progress-bar {
            height: 100%;
            background: linear-gradient(90deg, #00E5FF 0%, #2563EB 100%);
            border-radius: 9999px;
            animation: progressSweep 1.8s ease-in-out forwards;
        }

        .cute-pulsing-dot {
            width: 9px;
            height: 9px;
            background-color: #00E5FF;
            border-radius: 50%;
            animation: cuteBounceDot 0.8s ease-in-out infinite;
        }
    </style>
    <div id="splash-screen">
        <div class="splash-logo-container">
            <svg width="86" height="86" viewBox="0 0 110 100" fill="none">
                <polygon points="58,6 74,6 68,16 52,16" fill="#00E5FF" />
                <polygon points="35,16 55,16 35,50 44,50 41,58 30,58 12,90 2,90 22,58 17,58 20,50 26,50" fill="#0E1E38" />
                <polygon points="63,16 79,16 58,50 67,50 64,58 53,58 35,90 25,90 45,58 40,58 43,50 49,50" fill="#2563EB" />
            </svg>
        </div>
        <div class="splash-brand">Runway<span style="color:#2563EB;">x</span></div>
        <div class="splash-subtext">FINANCE CONTROLLER</div>
        <div class="splash-progress-wrapper">
            <div class="splash-progress-track">
                <div class="splash-progress-bar"></div>
            </div>
            <div class="cute-pulsing-dot"></div>
        </div>
    </div>
    """, unsafe_allow_html=True)

col_top_space, col_top_toggle = st.columns([5.5, 1.8])
with col_top_toggle:
    theme_mode = st.radio(
        "Theme Selector", 
        ["🌙 Dark", "☀️ Light"], 
        index=0, 
        horizontal=True, 
        label_visibility="collapsed"
    )
is_dark = "Dark" in theme_mode

if is_dark:
    st.markdown("""
    <style>
        div[data-baseweb="popover"] ul[data-testid="main-menu-list"] > li:first-child,
        div[data-baseweb="popover"] ul[data-testid="main-menu-list"] > hr:first-of-type,
        div[data-baseweb="popover"] [data-testid="stThemeSelector"],
        div[data-baseweb="popover"] div:has(> button[title="Light"]),
        div[data-baseweb="popover"] div:has(> button[title="Dark"]),
        div[data-baseweb="popover"] div:has(> button[aria-label="Light"]),
        div[data-baseweb="popover"] div:has(> button[aria-label="Dark"]),
        div[data-baseweb="popover"] div[role="radiogroup"] {
            display: none !important;
        }

        :root, [data-testid="stAppViewContainer"], .stApp {
            --text-color: #F8FAFC !important;
            --background-color: #020611 !important;
            --secondary-background-color: #060D1A !important;
            color: #F8FAFC !important;
        }

        html, body, .stApp, p, h1, h2, h3, h4, h5, h6, label, input, textarea {
            font-family: 'Plus Jakarta Sans', sans-serif !important;
        }

        [data-testid="stSidebarCollapseButton"] *,
        [data-testid="stSidebarCollapsedControl"] *,
        [data-testid="collapsedControl"] *,
        [data-testid="stIcon"],
        [data-testid="stIconMaterial"],
        [data-testid="stFileUploaderDropzone"] button [data-testid="stIcon"],
        [data-testid="stFileUploaderDropzone"] button [data-testid="stIconMaterial"],
        [data-testid="stFileUploaderDropzone"] button span:first-child,
        header button *,
        span[class*="material-symbols"],
        span[class*="material-icons"] {
            font-family: "Material Symbols Rounded", "Material Icons" !important;
            font-feature-settings: "liga" 1 !important;
            text-transform: none !important;
            letter-spacing: normal !important;
        }

        .stApp {
            background-color: #020611 !important;
            background-image: 
                radial-gradient(circle at 50% 0%, rgba(37, 99, 235, 0.08) 0%, transparent 45%),
                linear-gradient(180deg, #020611 0%, #040915 45%, #02050D 100%) !important;
            background-attachment: fixed !important;
            color: #F8FAFC !important;
        }

        header[data-testid="stHeader"] { background: transparent !important; }

        div[data-testid="stRadio"] > div[role="radiogroup"] {
            display: flex;
            justify-content: flex-end;
            gap: 6px;
            background: #060D1A;
            border: 1px solid #1E293B;
            border-radius: 9999px;
            padding: 4px 10px;
            width: fit-content;
            margin-left: auto;
            box-shadow: 0 4px 14px rgba(0, 0, 0, 0.5);
        }
        div[data-testid="stRadio"] label span {
            color: #F8FAFC !important;
            font-size: 12px !important;
            font-weight: 600 !important;
        }

        .hero-box { position: relative; text-align: center; padding: 0 0 24px 0; }
        
        .hero-badge-dark {
            display: inline-flex;
            align-items: center;
            gap: 6px;
            background: rgba(14, 30, 56, 0.9);
            border: 1px solid rgba(37, 99, 235, 0.5);
            color: #38BDF8;
            padding: 5px 16px;
            border-radius: 9999px;
            font-family: 'Space Grotesk', sans-serif !important;
            font-size: 11px !important;
            font-weight: 700 !important;
            letter-spacing: 1.2px !important;
            text-transform: uppercase;
            margin-bottom: 12px;
        }

        .main-title-dark {
            font-family: 'Space Grotesk', sans-serif !important;
            font-size: 40px !important;
            font-weight: 700 !important;
            letter-spacing: -1px !important;
            color: #FFFFFF !important;
            margin-bottom: 6px;
        }

        .title-accent-cyan {
            font-family: 'Space Grotesk', sans-serif !important;
            color: #38BDF8 !important;
        }

        .sub-title-dark {
            font-family: 'Plus Jakarta Sans', sans-serif !important;
            font-size: 14px !important;
            color: #94A3B8 !important;
            max-width: 740px;
            margin: 0 auto;
            line-height: 1.6 !important;
        }

        div[data-testid="stMetric"] {
            background-color: #060D1A !important;
            border: 1px solid #172338 !important;
            border-radius: 12px !important;
            padding: 16px 20px !important;
            box-shadow: 0 6px 18px rgba(0, 0, 0, 0.4) !important;
        }

        div[data-testid="stMetricLabel"], div[data-testid="stMetricLabel"] * {
            font-family: 'Plus Jakarta Sans', sans-serif !important;
            color: #94A3B8 !important;
            font-size: 11px !important;
            font-weight: 700 !important;
            letter-spacing: 0.8px !important;
            text-transform: uppercase;
        }

        div[data-testid="stMetricValue"], div[data-testid="stMetricValue"] * {
            color: #FFFFFF !important;
            font-family: 'Space Grotesk', sans-serif !important;
            font-weight: 700 !important;
            font-size: 26px !important;
            letter-spacing: -0.5px !important;
        }

        div[data-testid="stMetricDelta"] div {
            font-family: 'Plus Jakarta Sans', sans-serif !important;
            color: #38BDF8 !important;
            font-weight: 600 !important;
            font-size: 12px !important;
        }

        section[data-testid="stSidebar"] {
            background-color: #01040A !important;
            border-right: 1px solid #0F172A !important;
        }

        section[data-testid="stSidebar"] p,
        section[data-testid="stSidebar"] span,
        section[data-testid="stSidebar"] label,
        section[data-testid="stSidebar"] h2,
        section[data-testid="stSidebar"] h3 {
            color: #F1F5F9 !important;
        }

        input[type="checkbox"]:checked {
            accent-color: #EF4444 !important;
        }

        div.stButton > button {
            font-family: 'Space Grotesk', sans-serif !important;
            background: #2563EB !important;
            color: #FFFFFF !important;
            border-radius: 9999px !important;
            font-weight: 700 !important;
            font-size: 13px !important;
            border: none !important;
            padding: 8px 22px !important;
            box-shadow: 0 0 16px rgba(37, 99, 235, 0.35) !important;
            transition: all 0.2s ease !important;
        }

        div.stButton > button:hover {
            background: #1D4ED8 !important;
            box-shadow: 0 0 24px rgba(37, 99, 235, 0.5) !important;
            color: #FFFFFF !important;
        }

        div.stDownloadButton > button {
            font-family: 'Space Grotesk', sans-serif !important;
            background: #060D1A !important;
            border: 1px solid #2563EB !important;
            color: #38BDF8 !important;
            font-weight: 700 !important;
            border-radius: 9999px !important;
            padding: 9px 24px !important;
        }

        [data-testid="stFileUploader"] { background-color: transparent !important; }
        [data-testid="stFileUploader"] label, [data-testid="stFileUploader"] label p {
            color: #FFFFFF !important;
            font-family: 'Space Grotesk', sans-serif !important;
            font-weight: 700 !important;
            font-size: 14px !important;
        }
        [data-testid="stFileUploaderDropzone"], section[data-testid="stFileUploaderDropzone"] {
            background-color: #060D1A !important;
            border: 1.5px dashed #1E293B !important;
            border-radius: 14px !important;
            box-shadow: 0 4px 18px rgba(0, 0, 0, 0.4) !important;
            padding: 20px !important;
        }
        [data-testid="stFileUploaderDropzone"]:hover {
            border-color: #2563EB !important;
        }
        [data-testid="stFileUploaderDropzone"] span,
        [data-testid="stFileUploaderDropzone"] small,
        [data-testid="stFileUploaderDropzone"] p {
            color: #94A3B8 !important;
            font-family: 'Plus Jakarta Sans', sans-serif !important;
        }
        [data-testid="stFileUploaderDropzone"] button {
            background: #2563EB !important;
            color: #FFFFFF !important;
            border-radius: 9999px !important;
            border: none !important;
            font-weight: 700 !important;
            padding: 6px 20px !important;
            box-shadow: 0 0 12px rgba(37, 99, 235, 0.3) !important;
        }
        [data-testid="stFileUploaderDropzone"] button * { color: #FFFFFF !important; }
        [data-testid="stFileUploaderFile"] {
            background-color: #0B1322 !important;
            border: 1px solid #1E293B !important;
            border-radius: 10px !important;
        }
        [data-testid="stFileUploaderFile"] * { color: #FFFFFF !important; }

        details[data-testid="stExpander"] {
            background-color: #060D1A !important;
            border: 1px solid #1E293B !important;
            border-radius: 14px !important;
            box-shadow: 0 4px 18px rgba(0, 0, 0, 0.3) !important;
            margin-top: 14px !important;
        }
        details[data-testid="stExpander"] summary {
            background-color: #060D1A !important;
            border-radius: 14px !important;
            color: #FFFFFF !important;
            font-family: 'Space Grotesk', sans-serif !important;
            font-weight: 700 !important;
        }
        details[data-testid="stExpander"] summary * { color: #FFFFFF !important; font-weight: 700 !important; }
        details[data-testid="stExpander"] summary:hover { color: #38BDF8 !important; }
        details[data-testid="stExpander"] > div {
            background-color: #060D1A !important;
            border-top: 1px solid #172338 !important;
            padding: 16px !important;
        }

        div[data-baseweb="select"] > div {
            background-color: #081224 !important;
            border: 1px solid #1E293B !important;
            border-radius: 10px !important;
            color: #FFFFFF !important;
        }
        div[data-baseweb="select"] * { color: #FFFFFF !important; font-family: 'Plus Jakarta Sans', sans-serif !important; }
        div[data-baseweb="select"]:hover > div { border-color: #2563EB !important; }
        div[data-baseweb="popover"] ul[role="listbox"] {
            background-color: #081224 !important;
            border: 1px solid #1E293B !important;
            border-radius: 10px !important;
        }
        div[data-baseweb="popover"] li[role="option"] {
            background-color: #081224 !important;
            color: #FFFFFF !important;
        }
        div[data-baseweb="popover"] li[role="option"]:hover,
        div[data-baseweb="popover"] li[aria-selected="true"] {
            background-color: #172338 !important;
            color: #38BDF8 !important;
        }
        div[data-baseweb="popover"] li[role="option"] * { color: inherit !important; }

        div[data-testid="stSlider"] * { color: #F1F5F9 !important; }
        div[data-testid="stSlider"] div[data-baseweb="slider"] div { background-color: #2563EB !important; }
        div[data-testid="stAlert"] {
            background-color: #0A1424 !important;
            border: 1px solid #1E293B !important;
            border-radius: 12px !important;
            color: #F8FAFC !important;
        }
        div[data-testid="stAlert"] * { color: #F8FAFC !important; }

        div[data-baseweb="tab-list"] {
            background: transparent !important;
            border-bottom: 1px solid #1E293B !important;
            gap: 20px !important;
            padding: 0 !important;
        }
        div[data-baseweb="tab-list"] button[data-baseweb="tab"] {
            background: transparent !important;
            background-color: transparent !important;
            border: none !important;
            padding: 8px 10px !important;
            box-shadow: none !important;
        }
        div[data-baseweb="tab-list"] button[data-baseweb="tab"] * {
            background: transparent !important;
            background-color: transparent !important;
        }
        div[data-baseweb="tab-list"] button[data-baseweb="tab"] p {
            color: #94A3B8 !important;
            font-weight: 600 !important;
            font-size: 13px !important;
            margin: 0 !important;
        }
        div[data-baseweb="tab-list"] button[data-baseweb="tab"][aria-selected="true"] {
            border-bottom: 2.5px solid #2563EB !important;
        }
        div[data-baseweb="tab-list"] button[data-baseweb="tab"][aria-selected="true"] p {
            color: #38BDF8 !important;
            font-weight: 700 !important;
        }
        div[data-baseweb="tab-highlight"], div[data-baseweb="tab-border"] { background-color: transparent !important; }
    </style>
    """, unsafe_allow_html=True)

else:
    st.markdown("""
    <style>
        div[data-baseweb="popover"] ul[data-testid="main-menu-list"] > li:first-child,
        div[data-baseweb="popover"] ul[data-testid="main-menu-list"] > hr:first-of-type,
        div[data-baseweb="popover"] [data-testid="stThemeSelector"],
        div[data-baseweb="popover"] div:has(> button[title="Light"]),
        div[data-baseweb="popover"] div:has(> button[title="Dark"]),
        div[data-baseweb="popover"] div:has(> button[aria-label="Light"]),
        div[data-baseweb="popover"] div:has(> button[aria-label="Dark"]),
        div[data-baseweb="popover"] div[role="radiogroup"] {
            display: none !important;
        }

        :root, .stApp, [data-testid="stAppViewContainer"], [data-testid="stMarkdownContainer"] {
            --text-color: #0F172A !important;
            --primary-color: #0284C7 !important;
            --background-color: #EDF5FD !important;
            --secondary-background-color: #FFFFFF !important;
            color: #0F172A !important;
            font-family: 'Plus Jakarta Sans', sans-serif !important;
        }

        [data-testid="stSidebarCollapseButton"] *,
        [data-testid="stSidebarCollapsedControl"] *,
        [data-testid="collapsedControl"] *,
        [data-testid="stIcon"],
        [data-testid="stIconMaterial"],
        [data-testid="stFileUploaderDropzone"] button [data-testid="stIcon"],
        [data-testid="stFileUploaderDropzone"] button [data-testid="stIconMaterial"],
        [data-testid="stFileUploaderDropzone"] button span:first-child,
        header button *,
        span[class*="material-symbols"],
        span[class*="material-icons"] {
            font-family: "Material Symbols Rounded", "Material Icons" !important;
            font-feature-settings: "liga" 1 !important;
            text-transform: none !important;
            letter-spacing: normal !important;
        }

        .stApp {
            background-color: #EDF5FD !important;
            background-image: 
                linear-gradient(180deg, #99C5FA 0%, #C7E0FD 22%, #E8F3FE 55%, #F8FAFD 85%, #FFFFFF 100%) !important;
            background-attachment: fixed !important;
            color: #0F172A !important;
        }

        header[data-testid="stHeader"] { background: transparent !important; }

        div[data-testid="stRadio"] > div[role="radiogroup"] {
            display: flex;
            justify-content: flex-end;
            gap: 6px;
            background: #FFFFFF;
            border: 1px solid #E2E8F0;
            border-radius: 9999px;
            padding: 4px 10px;
            width: fit-content;
            margin-left: auto;
            box-shadow: 0 2px 8px rgba(0, 0, 0, 0.05);
        }
        div[data-testid="stRadio"] label span {
            color: #0F172A !important;
            font-size: 12px !important;
            font-weight: 600 !important;
        }

        .hero-box { position: relative; text-align: center; padding: 0 0 24px 0; }
        
        .hero-badge-light {
            display: inline-flex;
            align-items: center;
            gap: 8px;
            background: #FFFFFF;
            border: 1px solid #E2E8F0;
            box-shadow: 0 2px 10px rgba(2, 132, 199, 0.08);
            color: #0284C7;
            padding: 5px 16px;
            border-radius: 9999px;
            font-family: 'Space Grotesk', sans-serif !important;
            font-size: 11px !important;
            font-weight: 700 !important;
            letter-spacing: 1px !important;
            text-transform: uppercase;
            margin-bottom: 12px;
        }

        .hero-badge-dot-blue {
            width: 7px;
            height: 7px;
            background: #0284C7;
            border-radius: 2px;
        }

        .main-title-light {
            font-family: 'Space Grotesk', sans-serif !important;
            font-size: 40px !important;
            font-weight: 700 !important;
            letter-spacing: -1px !important;
            color: #0F172A !important;
            margin-bottom: 6px;
        }

        .title-accent-blue {
            font-family: 'Space Grotesk', sans-serif !important;
            color: #0284C7 !important;
        }

        .sub-title-light {
            font-family: 'Plus Jakarta Sans', sans-serif !important;
            font-size: 14px !important;
            color: #475569 !important;
            max-width: 740px;
            margin: 0 auto;
            line-height: 1.6 !important;
        }

        div[data-testid="stMetric"] {
            background-color: #FFFFFF !important;
            border: 1px solid #E2E8F0 !important;
            border-radius: 14px !important;
            padding: 16px 20px !important;
            box-shadow: 0 6px 20px rgba(30, 58, 138, 0.04) !important;
        }

        div[data-testid="stMetricLabel"], div[data-testid="stMetricLabel"] * {
            font-family: 'Plus Jakarta Sans', sans-serif !important;
            color: #64748B !important;
            font-size: 11px !important;
            font-weight: 700 !important;
            letter-spacing: 0.8px !important;
            text-transform: uppercase;
        }

        div[data-testid="stMetricValue"], div[data-testid="stMetricValue"] * {
            color: #0F172A !important;
            font-family: 'Space Grotesk', sans-serif !important;
            font-weight: 700 !important;
            font-size: 26px !important;
            letter-spacing: -0.5px !important;
        }

        div[data-testid="stMetricDelta"] div {
            font-family: 'Plus Jakarta Sans', sans-serif !important;
            color: #10B981 !important;
            font-weight: 600 !important;
            font-size: 12px !important;
        }

        section[data-testid="stSidebar"] {
            background: rgba(235, 244, 254, 0.75) !important;
            backdrop-filter: blur(25px) !important;
            border-right: 1px solid rgba(255, 255, 255, 0.8) !important;
        }

        section[data-testid="stSidebar"] p,
        section[data-testid="stSidebar"] span,
        section[data-testid="stSidebar"] label,
        section[data-testid="stSidebar"] h2,
        section[data-testid="stSidebar"] h3 {
            color: #1E293B !important;
        }

        input[type="checkbox"]:checked {
            accent-color: #EF4444 !important;
        }

        div.stButton > button {
            font-family: 'Space Grotesk', sans-serif !important;
            background: #0284C7 !important;
            color: #FFFFFF !important;
            border-radius: 9999px !important;
            font-weight: 600 !important;
            font-size: 13px !important;
            border: none !important;
            padding: 8px 22px !important;
            box-shadow: 0 4px 14px rgba(2, 132, 199, 0.3) !important;
            transition: all 0.2s ease !important;
        }

        div.stButton > button:hover {
            background: #0369A1 !important;
            color: #FFFFFF !important;
        }

        div.stDownloadButton > button {
            font-family: 'Space Grotesk', sans-serif !important;
            background: #FFFFFF !important;
            border: 1.5px solid #0284C7 !important;
            color: #0284C7 !important;
            font-weight: 700 !important;
            border-radius: 9999px !important;
            padding: 9px 24px !important;
        }

        [data-testid="stFileUploader"] { background-color: transparent !important; }
        [data-testid="stFileUploader"] label, [data-testid="stFileUploader"] label p {
            color: #0F172A !important;
            font-family: 'Space Grotesk', sans-serif !important;
            font-weight: 700 !important;
            font-size: 14px !important;
        }
        [data-testid="stFileUploaderDropzone"], section[data-testid="stFileUploaderDropzone"] {
            background-color: #FFFFFF !important;
            border: 1.5px dashed #93C5FD !important;
            border-radius: 14px !important;
            box-shadow: 0 4px 16px rgba(30, 58, 138, 0.04) !important;
            padding: 20px !important;
        }
        [data-testid="stFileUploaderDropzone"]:hover { border-color: #0284C7 !important; }
        [data-testid="stFileUploaderDropzone"] span,
        [data-testid="stFileUploaderDropzone"] small,
        [data-testid="stFileUploaderDropzone"] p {
            color: #475569 !important;
            font-family: 'Plus Jakarta Sans', sans-serif !important;
        }
        [data-testid="stFileUploaderDropzone"] button {
            background: #0284C7 !important;
            color: #FFFFFF !important;
            border-radius: 9999px !important;
            border: none !important;
            font-weight: 600 !important;
            padding: 6px 20px !important;
            box-shadow: 0 2px 10px rgba(2, 132, 199, 0.25) !important;
        }
        [data-testid="stFileUploaderDropzone"] button * { color: #FFFFFF !important; }
        [data-testid="stFileUploaderFile"] {
            background-color: #F8FAFC !important;
            border: 1px solid #E2E8F0 !important;
            border-radius: 10px !important;
        }
        [data-testid="stFileUploaderFile"] * { color: #0F172A !important; }

        details[data-testid="stExpander"] {
            background-color: #FFFFFF !important;
            border: 1px solid #E2E8F0 !important;
            border-radius: 14px !important;
            box-shadow: 0 4px 16px rgba(30, 58, 138, 0.03) !important;
            margin-top: 14px !important;
        }
        details[data-testid="stExpander"] summary {
            background-color: #FFFFFF !important;
            border-radius: 14px !important;
            color: #0F172A !important;
            font-family: 'Space Grotesk', sans-serif !important;
            font-weight: 700 !important;
        }
        details[data-testid="stExpander"] summary * { color: #0F172A !important; font-weight: 700 !important; }
        details[data-testid="stExpander"] summary:hover { color: #0284C7 !important; }
        details[data-testid="stExpander"] > div {
            background-color: #FFFFFF !important;
            border-top: 1px solid #F1F5F9 !important;
            padding: 16px !important;
        }

        div[data-baseweb="select"] > div {
            background-color: #FFFFFF !important;
            border: 1px solid #CBD5E1 !important;
            border-radius: 10px !important;
            color: #0F172A !important;
        }
        div[data-baseweb="select"] * { color: #0F172A !important; font-family: 'Plus Jakarta Sans', sans-serif !important; }
        div[data-baseweb="select"]:hover > div { border-color: #0284C7 !important; }
        div[data-baseweb="popover"] ul[role="listbox"] {
            background-color: #FFFFFF !important;
            border: 1px solid #E2E8F0 !important;
            border-radius: 10px !important;
            box-shadow: 0 8px 24px rgba(15, 23, 42, 0.08) !important;
        }
        div[data-baseweb="popover"] li[role="option"] { background-color: #FFFFFF !important; color: #0F172A !important; }
        div[data-baseweb="popover"] li[role="option"]:hover,
        div[data-baseweb="popover"] li[aria-selected="true"] { background-color: #F1F5F9 !important; color: #0284C7 !important; }
        div[data-baseweb="popover"] li[role="option"] * { color: inherit !important; }

        div[data-testid="stSlider"] * { color: #1E293B !important; }
        div[data-testid="stSlider"] div[data-baseweb="slider"] div { background-color: #0284C7 !important; }
        div[data-testid="stAlert"] {
            background-color: rgba(254, 240, 138, 0.5) !important;
            border: 1px solid #FACC15 !important;
            border-radius: 12px !important;
            color: #713F12 !important;
        }
        div[data-testid="stAlert"] * { color: #713F12 !important; }

        div[data-baseweb="tab-list"] {
            background: transparent !important;
            border-bottom: 1px solid #E2E8F0 !important;
            gap: 20px !important;
            padding: 0 !important;
        }
        div[data-baseweb="tab-list"] button[data-baseweb="tab"] {
            background: transparent !important;
            background-color: transparent !important;
            border: none !important;
            padding: 8px 10px !important;
            box-shadow: none !important;
        }
        div[data-baseweb="tab-list"] button[data-baseweb="tab"] * {
            background: transparent !important;
            background-color: transparent !important;
        }
        div[data-baseweb="tab-list"] button[data-baseweb="tab"] p {
            color: #64748B !important;
            font-weight: 600 !important;
            font-size: 13px !important;
            margin: 0 !important;
        }
        div[data-baseweb="tab-list"] button[data-baseweb="tab"][aria-selected="true"] {
            border-bottom: 2.5px solid #0284C7 !important;
        }
        div[data-baseweb="tab-list"] button[data-baseweb="tab"][aria-selected="true"] p {
            color: #0284C7 !important;
            font-weight: 700 !important;
        }
        div[data-baseweb="tab-highlight"], div[data-baseweb="tab-border"] { background-color: transparent !important; }
    </style>
    """, unsafe_allow_html=True)

logo_text_color = "#FFFFFF" if is_dark else "#0F172A"
sub_brand_color = "#94A3B8" if is_dark else "#64748B"

st.sidebar.markdown(f"""
<div style="display: flex; align-items: center; gap: 12px; margin-bottom: 26px; padding: 4px 0;">
    <svg width="46" height="42" viewBox="0 0 110 100" fill="none" style="flex-shrink: 0;">
        <polygon points="58,6 74,6 68,16 52,16" fill="#00E5FF" />
        <polygon points="35,16 55,16 35,50 44,50 41,58 30,58 12,90 2,90 22,58 17,58 20,50 26,50" fill="{'#FFFFFF' if is_dark else '#0E1E38'}" />
        <polygon points="63,16 79,16 58,50 67,50 64,58 53,58 35,90 25,90 45,58 40,58 43,50 49,50" fill="#2563EB" />
    </svg>
    <div>
        <div style="font-family: 'Space Grotesk', sans-serif; font-weight: 800; font-size: 22px; line-height: 1; font-style: italic; color: {logo_text_color};">
            Runway<span style="color:#2563EB;">x</span>
        </div>
        <div style="font-family: 'Plus Jakarta Sans', sans-serif; font-size: 8px; font-weight: 700; letter-spacing: 2px; color: {sub_brand_color}; margin-top: 4px; text-transform: uppercase;">
            FINANCE CONTROLLER
        </div>
    </div>
</div>
""", unsafe_allow_html=True)

def render_styled_table(df, dark_mode=True):
    if df.empty:
        return "<div style='color: #94A3B8; padding: 20px; text-align: center;'>No records found.</div>"

    if dark_mode:
        container_style = (
            "width: 100%; max-height: 520px; overflow-y: auto; overflow-x: auto; "
            "border: 1px solid #172338; border-radius: 12px; background-color: #060D1A; "
            "box-shadow: 0 8px 24px rgba(0, 0, 0, 0.4);"
        )
        th_style = (
            "background-color: #0B1322; color: #94A3B8; font-family: 'Space Grotesk', sans-serif; "
            "font-weight: 700; font-size: 11px; padding: 12px 16px; border-bottom: 1px solid #1E293B; "
            "position: sticky; top: 0; z-index: 10; white-space: nowrap; text-align: center; letter-spacing: 0.5px;"
        )
        td_base = (
            "background-color: #060D1A; color: #FFFFFF !important; font-family: 'Plus Jakarta Sans', sans-serif; "
            "font-size: 12.5px; font-weight: 500; padding: 11px 16px; border-bottom: 1px solid #172338; "
            "white-space: nowrap; text-align: center;"
        )
    else:
        container_style = (
            "width: 100%; max-height: 520px; overflow-y: auto; overflow-x: auto; "
            "border: 1px solid #E2E8F0; border-radius: 12px; background-color: #FFFFFF; "
            "box-shadow: 0 4px 16px rgba(30, 58, 138, 0.03);"
        )
        th_style = (
            "background-color: #F8FAFC; color: #475569; font-family: 'Space Grotesk', sans-serif; "
            "font-weight: 700; font-size: 11px; padding: 12px 16px; border-bottom: 1px solid #E2E8F0; "
            "position: sticky; top: 0; z-index: 10; white-space: nowrap; text-align: center; letter-spacing: 0.5px;"
        )
        td_base = (
            "background-color: #FFFFFF; color: #0F172A !important; font-family: 'Plus Jakarta Sans', sans-serif; "
            "font-size: 12.5px; font-weight: 500; padding: 11px 16px; border-bottom: 1px solid #E2E8F0; "
            "white-space: nowrap; text-align: center;"
        )

    headers = "".join([f"<th style='{th_style}'>{col}</th>" for col in df.columns])
    
    rows = []
    for _, row in df.iterrows():
        cells = []
        for col, val in row.items():
            val_str = str(val) if val is not None else ""
            curr_td = td_base
            
            if dark_mode:
                if col in ['Order ID', 'Bank UTR', 'Bank UTR Ref']:
                    curr_td += " color: #38BDF8 !important;"
                elif col in ['AI Confidence']:
                    curr_td += " color: #00F298 !important; font-weight: 700;"
            
            cells.append(f"<td style='{curr_td}'>{val_str}</td>")
        rows.append(f"<tr>{''.join(cells)}</tr>")

    table_html = f"""
    <div style="{container_style}">
        <table style="width: 100%; border-collapse: collapse; margin: 0;">
            <thead>
                <tr>{headers}</tr>
            </thead>
            <tbody>
                {''.join(rows)}
            </tbody>
        </table>
    </div>
    """
    return table_html

if is_dark:
    st.markdown("""
    <div class="hero-box">
        <div class="hero-badge-dark">
            ⚡ RUNWAYX AUTONOMOUS SETTLEMENT
        </div>
        <div class="main-title-dark">Reconciliation <span class="title-accent-cyan">Secure & Faster</span></div>
        <div class="sub-title-dark">Multi-source deterministic settlement, MDR/GST tax variance audit, and automated gateway remediation.</div>
    </div>
    """, unsafe_allow_html=True)
else:
    st.markdown("""
    <div class="hero-box">
        <div class="hero-badge-light">
            <span class="hero-badge-dot-blue"></span>
            RUNWAYX SETTLEMENT ARCHITECTURE
        </div>
        <div class="main-title-light">Smarter Settlements. <span class="title-accent-blue">Zero Friction.</span></div>
        <div class="sub-title-light">Deterministic matching engine, algorithmic tax/fee deduction auditing, and instant dispute dispatch.</div>
    </div>
    """, unsafe_allow_html=True)

def find_best_col_idx(columns, keywords):
    for kw in keywords:
        for idx, col in enumerate(columns):
            if kw.lower() in str(col).lower():
                return idx
    return 0

st.sidebar.header("⚙️ Data Configuration")
use_sample = st.sidebar.checkbox("Use Demo Synthetic Data (80 Records)", value=True)

ledger_df = None
bank_df = None
l_map = {}
b_map = {}

if use_sample:
    if not os.path.exists("internal_ledger.csv") or not os.path.exists("bank_statement.csv"):
        generate_synthetic_data(80)
    ledger_df = pd.read_csv("internal_ledger.csv")
    bank_df = pd.read_csv("bank_statement.csv")
    
    l_map = {
        'ref_id': 'order_id',
        'customer': 'customer_name',
        'amount': 'ledger_amount',
        'date': 'ledger_date'
    }
    b_map = {
        'bank_ref': 'utr_number',
        'narration': 'narration',
        'amount': 'bank_amount',
        'date': 'bank_date'
    }
    
    if st.sidebar.button("🔄 Regenerate Fresh Demo Records"):
        generate_synthetic_data(80)
        st.sidebar.success("Generated 80 fresh records!")
        st.rerun()
else:
    st.sidebar.subheader("📂 Custom Data Upload")
    max_rows = st.sidebar.slider("Maximum Rows to Ingest (Performance Cap)", 1000, 50000, 10000, step=1000)
    
    col_u1, col_u2 = st.columns(2)
    with col_u1:
        ledger_file = st.file_uploader("Upload Internal Ledger CSV", type=["csv"])
        if ledger_file:
            ledger_df = pd.read_csv(ledger_file, nrows=max_rows)
            ledger_df.columns = ledger_df.columns.astype(str).str.strip()
            st.success(f"Loaded {len(ledger_df):,} ledger records.")
            
            with st.expander("🛠️ Review / Map Ledger Columns", expanded=True):
                l_cols = ledger_df.columns.tolist()
                ref_idx = find_best_col_idx(l_cols, ['payment_id', 'order_id', 'ref', 'invoice', 'id'])
                l_map['ref_id'] = st.selectbox("Reference / Order ID Column", l_cols, index=ref_idx)
                
                cust_idx = find_best_col_idx(l_cols, ['customer', 'loan_id', 'vendor', 'party', 'account'])
                l_map['customer'] = st.selectbox("Customer / Vendor Column", ["None"] + l_cols, index=cust_idx + 1 if cust_idx < len(l_cols) else 0)
                if l_map['customer'] == "None": 
                    l_map['customer'] = None
                
                if 'Debit' in l_cols and 'Credit' in l_cols:
                    st.info("Mapped to dual Debit / Credit columns.")
                else:
                    amt_idx = find_best_col_idx(l_cols, ['amount_paid', 'ledger_amount', 'amount', 'amt', 'credit'])
                    l_map['amount'] = st.selectbox("Amount Column", l_cols, index=amt_idx)
                    
                date_idx = find_best_col_idx(l_cols, ['payment_date', 'ledger_date', 'date', 'time', 'dt'])
                l_map['date'] = st.selectbox("Ledger Date Column", l_cols, index=date_idx)

    with col_u2:
        bank_file = st.file_uploader("Upload Bank Statement CSV", type=["csv"])
        if bank_file:
            bank_df = pd.read_csv(bank_file, nrows=max_rows)
            bank_df.columns = bank_df.columns.astype(str).str.strip()
            st.success(f"Loaded {len(bank_df):,} bank records.")
            
            with st.expander("🛠️ Review / Map Bank Columns", expanded=True):
                b_cols = bank_df.columns.tolist()
                utr_idx = find_best_col_idx(b_cols, ['transaction_id', 'utr', 'chq', 'ref', 'id'])
                b_map['bank_ref'] = st.selectbox("UTR / Chq No Column", ["Auto-Generate"] + b_cols, index=utr_idx + 1 if utr_idx < len(b_cols) else 0)
                if b_map['bank_ref'] == "Auto-Generate": 
                    b_map['bank_ref'] = None
                
                narr_idx = find_best_col_idx(b_cols, ['channel', 'narration', 'transaction details', 'merchant_category', 'desc', 'type'])
                b_map['narration'] = st.selectbox("Narration Column", b_cols, index=narr_idx)
                
                if 'DEPOSIT AMT' in b_cols or 'WITHDRAWAL AMT' in b_cols:
                    st.info("Mapped to separate Deposit / Withdrawal columns.")
                else:
                    amt_idx = find_best_col_idx(b_cols, ['bank_amount', 'deposit', 'amount', 'amt'])
                    b_map['amount'] = st.selectbox("Bank Amount Column", b_cols, index=amt_idx)
                
                date_idx = find_best_col_idx(b_cols, ['txn_date', 'bank_date', 'date', 'time', 'dt'])
                b_map['date'] = st.selectbox("Bank Date Column", b_cols, index=date_idx)

if ledger_df is not None and bank_df is not None:
    try:
        agent = AIFinanceController(ledger_df, bank_df, l_map, b_map)
    except TypeError:
        agent = AIFinanceController(ledger_df, bank_df)
        
    reconciled_df, exceptions_df, summary = agent.run()

    total_l = summary.get('total_ledger') or summary.get('total_ledger_records', 0)
    total_b = summary.get('total_bank') or summary.get('total_bank_records', 0)
    rec_cnt = summary.get('reconciled_records', 0)
    m_rate = summary.get('match_rate_pct', 0.0)
    exc_cnt = summary.get('exceptions_count', 0)
    rec_vol = summary.get('reconciled_vol') or (reconciled_df['ledger_amount'].sum() if not reconciled_df.empty else 0.0)
    tax_var = summary.get('fee_variance') or summary.get('total_fee_variance', 0.0)

    k1, k2, k3, k4, k5, k6 = st.columns(6)
    k1.metric("Ledger Records", f"{total_l:,}")
    k2.metric("Bank Records", f"{total_b:,}")
    k3.metric("Match Rate", f"{m_rate}%", f"↑ {rec_cnt} Settled")
    k4.metric("Exceptions", f"{exc_cnt:,}", delta_color="inverse")
    k5.metric("Reconciled Volume", f"₹{rec_vol:,.2f}")
    k6.metric("Audited MDR/Tax", f"₹{tax_var:,.2f}")

    st.markdown("<div style='margin-bottom: 22px;'></div>", unsafe_allow_html=True)

    if is_dark:
        tab_names = [
            "Reconciled Records & Audit Trail", 
            "Exceptions & AI Dispute Dispatch", 
            "Analytics & Visualizations", 
            "Certified Audit Export"
        ]
    else:
        tab_names = [
            "Reconciled Records & Trail", 
            "Anomalies & AI Dispute Hub", 
            "Analytics & Visualizations", 
            "Certified Audit Export"
        ]

    tab1, tab2, tab3, tab4 = st.tabs(tab_names)

    with tab1:
        tab1_heading = "✅ Reconciled Records & Verified Audit Trail" if is_dark else "✅ Verified Settlement Audit Trail"
        header_color = "#FFFFFF" if is_dark else "#0F172A"
        st.markdown(f"<h4 style='color: {header_color}; font-family: Space Grotesk, sans-serif; font-weight: 700; margin-bottom: 14px;'>{tab1_heading}</h4>", unsafe_allow_html=True)
        
        if not reconciled_df.empty:
            table_display = reconciled_df.copy()
            table_display.insert(0, 'S.No', range(1, len(table_display) + 1))
            
            table_display['ledger_amount'] = table_display['ledger_amount'].apply(lambda x: f"₹{x:,.2f}")
            table_display['bank_amount'] = table_display['bank_amount'].apply(lambda x: f"₹{x:,.2f}")
            table_display['fee_deducted'] = table_display['fee_deducted'].apply(lambda x: f"₹{x:,.2f}")
            table_display['confidence_score'] = table_display['confidence_score'].apply(lambda x: f"{int(x * 100)}%")

            if is_dark:
                table_display.rename(columns={
                    'order_id': 'Order ID',
                    'invoice_id': 'Invoice Ref',
                    'customer_name': 'Customer / Party',
                    'ledger_amount': 'Ledger Amount',
                    'bank_amount': 'Bank Settlement',
                    'fee_deducted': 'MDR + GST Fee',
                    'utr_number': 'Bank UTR',
                    'match_type': 'Audit Category',
                    'confidence_score': 'AI Confidence',
                    'audit_notes': 'Auditor Remarks'
                }, inplace=True)
            else:
                table_display.rename(columns={
                    'order_id': 'Order ID',
                    'invoice_id': 'Invoice Ref',
                    'customer_name': 'Customer / Counterparty',
                    'ledger_amount': 'Ledger Amount',
                    'bank_amount': 'Bank Net Settlement',
                    'fee_deducted': 'MDR + GST Variance',
                    'utr_number': 'Bank UTR Ref',
                    'match_type': 'Match Strategy',
                    'confidence_score': 'AI Confidence',
                    'audit_notes': 'Auditor Remarks'
                }, inplace=True)

            st.markdown(render_styled_table(table_display, dark_mode=is_dark), unsafe_allow_html=True)
        else:
            st.info("No matching settlement pairs discovered across uploaded sources.")

    with tab2:
        st.markdown(f"<h4 style='color: {header_color}; font-family: Space Grotesk, sans-serif; font-weight: 700; margin-bottom: 14px;'>🚨 Unresolved Discrepancies</h4>", unsafe_allow_html=True)
        if not exceptions_df.empty:
            st.markdown(render_styled_table(exceptions_df, dark_mode=is_dark), unsafe_allow_html=True)
            
            missing_ledger = exceptions_df[exceptions_df['source'] == 'INTERNAL_LEDGER']
            if not missing_ledger.empty:
                action_color = "#38BDF8" if is_dark else "#0284C7"
                st.markdown(f"<h5 style='color: {action_color}; font-family: Space Grotesk, sans-serif; margin-top: 24px; font-weight: 700;'>✉️ 1-Click Autonomous Dispute Dispatch</h5>", unsafe_allow_html=True)
                selected_ref = st.selectbox("Target Exception ID:", missing_ledger['reference_id'].tolist())
                sel_row = missing_ledger[missing_ledger['reference_id'] == selected_ref].iloc[0]
                
                draft_ticket = AIFinanceController.generate_ai_dispute_email(
                    sel_row['reference_id'],
                    sel_row['customer_or_sender'],
                    sel_row['amount'],
                    sel_row['date']
                )
                st.text_area("Generated Formal Gateway Escalation Notice", draft_ticket, height=200)
                if st.button("🚀 Transmit Escalation to Gateway Desk"):
                    st.success(f"Dispute ticket dispatched to Razorpay Gateway Desk for reference {selected_ref}!")

    with tab3:
        col_c1, col_c2 = st.columns(2)
        
        text_color = "#FFFFFF" if is_dark else "#0F172A"
        label_color = "#E2E8F0" if is_dark else "#475569"
        bg_color = "#060D1A" if is_dark else "#FFFFFF"
        grid_color = "#1E293B" if is_dark else "#F1F5F9"
        
        chart_palette = ["#2563EB", "#00E5FF", "#38BDF8", "#FBBF24"] if is_dark else ["#0284C7", "#2563EB", "#10B981", "#F59E0B"]
        bar_palette = ["#2563EB", "#00E5FF"] if is_dark else ["#0284C7", "#2563EB"]
        
        with col_c1:
            if not reconciled_df.empty:
                fig1 = px.pie(
                    reconciled_df, 
                    names='match_type', 
                    title="Settlement Execution Breakdown", 
                    color_discrete_sequence=chart_palette,
                    hole=0.62
                )
                fig1.update_traces(
                    textfont=dict(color="#FFFFFF", size=13, family="Space Grotesk")
                )
                fig1.update_layout(
                    paper_bgcolor=bg_color,
                    plot_bgcolor=bg_color,
                    title=dict(font=dict(size=15, color=text_color, family="Space Grotesk")),
                    legend=dict(
                        font=dict(color=label_color, size=12, family="Plus Jakarta Sans")
                    ),
                    margin=dict(t=40, b=20, l=10, r=10)
                )
                st.plotly_chart(fig1, use_container_width=True)
                
        with col_c2:
            if not exceptions_df.empty:
                fig2 = px.bar(
                    exceptions_df, 
                    x='exception_category', 
                    color='source', 
                    title="Discrepancies by Source Channel", 
                    color_discrete_sequence=bar_palette
                )
                fig2.update_layout(
                    paper_bgcolor=bg_color,
                    plot_bgcolor=bg_color,
                    title=dict(font=dict(size=15, color=text_color, family="Space Grotesk")),
                    legend=dict(
                        title=dict(font=dict(color=text_color, size=12, family="Plus Jakarta Sans")),
                        font=dict(color=label_color, size=12, family="Plus Jakarta Sans")
                    ),
                    xaxis=dict(
                        title=dict(font=dict(color=label_color, size=12, family="Plus Jakarta Sans")),
                        tickfont=dict(color=label_color, size=11, family="Plus Jakarta Sans"),
                        gridcolor=grid_color,
                        linecolor=grid_color
                    ),
                    yaxis=dict(
                        title=dict(font=dict(color=label_color, size=12, family="Plus Jakarta Sans")),
                        tickfont=dict(color=label_color, size=11, family="Plus Jakarta Sans"),
                        gridcolor=grid_color,
                        linecolor=grid_color
                    ),
                    margin=dict(t=40, b=20, l=10, r=10)
                )
                st.plotly_chart(fig2, use_container_width=True)

    with tab4:
        st.markdown(f"<h4 style='color: {header_color}; font-family: Space Grotesk, sans-serif; font-weight: 700; margin-bottom: 14px;'>📑 Statutory Financial Compliance Package</h4>", unsafe_allow_html=True)
        buffer = io.BytesIO()
        
        with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
            reconciled_df.to_excel(writer, sheet_name="Reconciled_Trail", index=False)
            exceptions_df.to_excel(writer, sheet_name="Exceptions_List", index=False)
            
            summary_display = pd.DataFrame([
                {"Audit Metric": "Total Internal Ledger Orders", "Value": summary.get('total_ledger_records', 0)},
                {"Audit Metric": "Total Bank Credits Received", "Value": summary.get('total_bank_records', 0)},
                {"Audit Metric": "Successfully Reconciled Count", "Value": summary.get('reconciled_records', 0)},
                {"Audit Metric": "Reconciliation Match Rate (%)", "Value": f"{summary.get('match_rate_pct', 0)}%"},
                {"Audit Metric": "Unresolved Exceptions Count", "Value": summary.get('exceptions_count', 0)},
                {"Audit Metric": "Total Reconciled Volume (₹)", "Value": f"₹{summary.get('reconciled_vol', 0):,.2f}"},
                {"Audit Metric": "Audited Gateway MDR / Tax (₹)", "Value": f"₹{summary.get('fee_variance', 0):,.2f}"}
            ])
            summary_display.to_excel(writer, sheet_name="Executive_Summary", index=False)
            
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter

            for sheetname in writer.sheets:
                ws = writer.sheets[sheetname]
                for cell in ws[1]:
                    header_fill = "040A18" if is_dark else "0284C7"
                    header_font = "38BDF8" if is_dark else "FFFFFF"
                    cell.fill = PatternFill(start_color=header_fill, end_color=header_fill, fill_type="solid")
                    cell.font = Font(color=header_font, bold=True)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                
                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    st.download_button(
        "📥 Download Certified Audit Package (.xlsx)",
        data=buffer.getvalue(),
        file_name=f"RunwayX_Reconciliation_Audit_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

else:
    st.warning("Please upload both CSV datasets or check 'Use Demo Synthetic Data' to start the controller.")

footer_border = "#172338" if is_dark else "#E2E8F0"
footer_bg = "rgba(6, 13, 26, 0.7)" if is_dark else "rgba(255, 255, 255, 0.7)"
footer_text = "#64748B" if is_dark else "#64748B"
primary_brand = "#38BDF8" if is_dark else "#0284C7"

st.markdown(f"""
<div style="margin-top: 50px; padding: 22px 16px 16px; border-top: 1px solid {footer_border}; background: {footer_bg}; border-radius: 12px; text-align: center;">
    <div style="display: flex; justify-content: center; align-items: center; gap: 8px; margin-bottom: 6px;">
        <span style="font-family: 'Space Grotesk', sans-serif; font-size: 13px; font-weight: 700; color: {primary_brand};">
            Powered by Razorpay Settlement & Discrepancy Gateway Infrastructure
        </span>
    </div>
    <div style="font-family: 'Plus Jakarta Sans', sans-serif; font-size: 11px; font-weight: 600; letter-spacing: 0.4px; color: {footer_text};">
        Designed & Developed by <span style="font-weight: 700; color: {'#F8FAFC' if is_dark else '#0F172A'};">Dhruv Tomar</span>. All Rights Reserved © {time.strftime('%Y')}
    </div>
</div>
""", unsafe_allow_html=True)